"""Grille de relecture par des enseignants — phase 4 du plan d'entraînement.

Le banc automatique juge la forme ; la VÉRITÉ du contenu (une date d'histoire,
un énoncé sans question, une explication fausse) ne se juge qu'à la lecture.

    python recherche/src/relecture.py preparer
    python recherche/src/relecture.py analyser recherche/relecture/retours/*.xlsx

`preparer` tire, à graine fixe, des exercices et des corrections produits par
le modèle affiné sur le banc, et y MÊLE sans le dire des témoins produits par
Gemini : les enseignants notent à l'aveugle, et leurs notes du modèle se
lisent en regard de celles de la référence — même relecteur, même jour, même
exigence. La clé (code → modèle) reste dans le dépôt et n'est pas envoyée.

`analyser` lit les grilles remplies, lève l'anonymat avec la clé, et calcule
par source : note moyenne, part d'exercices justes, part utilisable en
classe, par niveau ; et, si plusieurs enseignants ont noté les mêmes codes,
leur accord.
"""
from __future__ import annotations

import argparse
import collections
import json
import pathlib
import random
import re
import statistics
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import banc  # noqa: E402

RACINE = pathlib.Path(__file__).resolve().parents[2]
DOSSIER = RACINE / "recherche/relecture"
GRILLE = DOSSIER / "grille_relecture_v1.xlsx"
CLE = DOSSIER / "cle_relecture_v1.json"
RAPPORT = DOSSIER / "rapport_relecture_v1.json"

MODELE = "affine:Qwen/Qwen3.5-4B+repetia-v1"
TEMOIN = "gemini:gemini-flash-lite-latest"
EFFECTIFS = {"generation": 50, "correction": 50}
TEMOINS = {"generation": 10, "correction": 10}
GRAINE = 20260926
NIVEAUX = ["6ème", "5ème", "4ème", "BEPC", "BAC"]
CLASSE = {"BEPC": "3ème (BEPC)", "BAC": "Terminale (BAC)"}

# Colonnes à remplir, avec leurs valeurs permises (listes déroulantes).
NOTE = ["1", "2", "3", "4", "5"]
JUSTESSE = ["Juste", "Erreur mineure", "Erreur grave"]
UTILISABLE = ["Tel quel", "Après retouche", "Non"]
CRITERES = {
    "generation": [
        ("Justesse (solution et explication)", JUSTESSE),
        ("Respecte le thème demandé", ["Oui", "En partie", "Non"]),
        ("Adapté au niveau", ["Trop facile", "Adapté", "Trop difficile"]),
        ("Clarté de l'énoncé (1 à 5)", NOTE),
        ("Qualité de l'explication (1 à 5)", NOTE),
        ("Qualité du français (1 à 5)", NOTE),
        ("Utilisable en classe", UTILISABLE),
        ("Commentaire (facultatif)", None),
    ],
    "correction": [
        ("Le verdict (juste / faux) est-il le bon ?", ["Oui", "Non"]),
        ("Justesse de l'explication", JUSTESSE),
        ("Ton bienveillant (1 à 5)", NOTE),
        ("Utile à l'élève (1 à 5)", NOTE),
        ("Qualité du français (1 à 5)", NOTE),
        ("Utilisable avec un élève", UTILISABLE),
        ("Commentaire (facultatif)", None),
    ],
}

# ---------------------------------------------------------------------------
# Tirage
# ---------------------------------------------------------------------------


def _enonce_de_correction(it: dict) -> str:
    m = re.search(r"Exercice : (.*?)\. Solution attendue : (.*?)\. Réponse de l'élève : ", it["consigne"], re.S)
    return m.group(1) if m else ""


def _solution_de_correction(it: dict) -> str:
    m = re.search(r"Solution attendue : (.*?)\. Réponse de l'élève : ", it["consigne"], re.S)
    return m.group(1) if m else ""


def _eligibles(modele: str, tache: str, reponses: dict, items: dict, exclure: set) -> list[tuple[dict, dict]]:
    """Réponses conformes au schéma (une réponse illisible ne se relit pas :
    elle est déjà comptée comme défaut par le banc)."""
    sortie = []
    for id_, e in reponses.get(modele, {}).items():
        it = items[id_]
        if it["tache"] != tache or id_ in exclure:
            continue
        obj = banc.extraire_json(e.get("texte", ""))
        if banc.conforme(obj, tache):
            sortie.append((it, obj))
    return sortie


def _stratifie(candidats: list, n: int, alea: random.Random) -> list:
    """Tour à tour sur les niveaux, puis sur les matières dans chaque niveau."""
    par_niveau: dict[str, dict[str, list]] = collections.defaultdict(lambda: collections.defaultdict(list))
    for c in candidats:
        par_niveau[c[0]["niveau"]][c[0]["matiere"]].append(c)
    for mats in par_niveau.values():
        for liste in mats.values():
            alea.shuffle(liste)
    retenus = []
    rang = {niv: 0 for niv in par_niveau}
    while len(retenus) < n and any(any(m.values()) for m in par_niveau.values()):
        for niv in [x for x in NIVEAUX if x in par_niveau]:
            mats = [m for m in sorted(par_niveau[niv]) if par_niveau[niv][m]]
            if not mats or len(retenus) >= n:
                continue
            m = mats[rang[niv] % len(mats)]
            rang[niv] += 1
            retenus.append(par_niveau[niv][m].pop())
    return retenus


def tirer() -> list[dict]:
    items = {i["id"]: i for i in banc.charger_jeu()}
    reponses = banc.reponses_par_modele()
    alea = random.Random(GRAINE)
    tirage = []
    for tache in ("generation", "correction"):
        temoins = _stratifie(_eligibles(TEMOIN, tache, reponses, items, set()), TEMOINS[tache], alea)
        # Le modèle et les témoins portent sur des items DIFFÉRENTS : un même
        # énoncé noté deux fois trahirait la présence d'un témoin.
        exclus = {it["id"] for it, _ in temoins}
        principaux = _stratifie(_eligibles(MODELE, tache, reponses, items, exclus), EFFECTIFS[tache], alea)
        for source, lot in ((MODELE, principaux), (TEMOIN, temoins)):
            for it, obj in lot:
                tirage.append({"tache": tache, "source": source, "item": it, "reponse": obj})
    alea.shuffle(tirage)
    compteurs = collections.Counter()
    for t in tirage:
        prefixe = "E" if t["tache"] == "generation" else "C"
        compteurs[prefixe] += 1
        t["code"] = f"{prefixe}{compteurs[prefixe]:03d}"
    return sorted(tirage, key=lambda t: t["code"])


# ---------------------------------------------------------------------------
# Classeur
# ---------------------------------------------------------------------------


def preparer() -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    tirage = tirer()
    DOSSIER.mkdir(parents=True, exist_ok=True)
    police = "Arial"
    titre = Font(name=police, size=14, bold=True, color="0F5F52")
    gras = Font(name=police, bold=True, color="FFFFFF")
    normal = Font(name=police, size=10)
    exemple = Font(name=police, size=10, italic=True, color="7F7F7F")
    fond_entete = PatternFill("solid", fgColor="0F5F52")
    fond_saisie = PatternFill("solid", fgColor="FFF2CC")  # jaune clair : à remplir
    fond_exemple = PatternFill("solid", fgColor="F2F2F2")
    trait = Side(style="thin", color="BFBFBF")
    bord = Border(left=trait, right=trait, top=trait, bottom=trait)
    haut = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws = wb.active
    ws.title = "Consignes"
    lignes = [
        ("Relecture des exercices et corrections de RépétIA", titre),
        ("", None),
        ("Merci de votre aide. Vous allez lire des exercices et des corrections rédigés automatiquement "
         "pour des élèves béninois, de la 6ème à la Terminale, et dire s'ils sont justes et utilisables.", normal),
        ("", None),
        ("Vos informations", Font(name=police, bold=True)),
        ("Nom et prénom :", normal),
        ("Discipline(s) enseignée(s) :", normal),
        ("Classes enseignées :", normal),
        ("Établissement (facultatif) :", normal),
        ("", None),
        ("Comment remplir", Font(name=police, bold=True)),
        ("1. Onglets « Exercices » et « Corrections » : ne remplissez QUE les cellules jaunes.", normal),
        ("2. Chaque cellule jaune propose une liste déroulante : choisissez une valeur.", normal),
        ("3. Relisez seulement les lignes de VOTRE discipline et de VOS niveaux (filtrez les colonnes "
         "« Matière » et « Niveau ») ; laissez les autres vides.", normal),
        ("4. La ligne grisée « EXEMPLE » montre le format attendu ; elle n'est pas comptée.", normal),
        ("5. Justesse : « Erreur mineure » = imprécision qui ne trompe pas l'élève ; « Erreur grave » = "
         "l'élève apprendrait quelque chose de faux (calcul, date, définition…).", normal),
        ("6. Notes de 1 (très mauvais) à 5 (excellent). Notez comme pour le travail d'un collègue.", normal),
        ("7. Un commentaire court aide beaucoup quand vous signalez une erreur : dites laquelle.", normal),
        ("", None),
        ("Les lignes ne sont pas toutes produites par le même système : ne cherchez pas à les distinguer, "
         "jugez chacune pour elle-même.", exemple),
        ("", None),
        ("Renvoyez le fichier rempli, renommé avec votre nom (ex. grille_relecture_v1_Dossou.xlsx).", normal),
    ]
    for r, (texte, f) in enumerate(lignes, 1):
        c = ws.cell(row=r, column=1, value=texte)
        if f:
            c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in (6, 7, 8, 9):
        ws.cell(row=r, column=2).fill = fond_saisie
        ws.cell(row=r, column=2).border = bord
    ws.column_dimensions["A"].width = 100
    ws.column_dimensions["B"].width = 40

    def feuille(nom: str, tache: str, colonnes_info: list[tuple[str, int]], lignes_info: list[list[str]],
                exemple_info: list[str], exemple_eval: list[str]) -> None:
        f = wb.create_sheet(nom)
        criteres = CRITERES[tache]
        entetes = [c for c, _ in colonnes_info] + [c for c, _ in criteres]
        for j, h in enumerate(entetes, 1):
            c = f.cell(row=1, column=j, value=h)
            c.font, c.fill, c.border = gras, fond_entete, bord
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for j, (_, largeur) in enumerate(colonnes_info, 1):
            f.column_dimensions[get_column_letter(j)].width = largeur
        debut_eval = len(colonnes_info) + 1
        for k, (crit, _) in enumerate(criteres):
            f.column_dimensions[get_column_letter(debut_eval + k)].width = 40 if crit.startswith("Commentaire") else 17
        # Ligne d'exemple, grisée.
        for j, v in enumerate(exemple_info + exemple_eval, 1):
            c = f.cell(row=2, column=j, value=v)
            c.font, c.fill, c.border, c.alignment = exemple, fond_exemple, bord, haut
        derniere = 2 + len(lignes_info)
        for i, valeurs in enumerate(lignes_info, 3):
            for j, v in enumerate(valeurs, 1):
                c = f.cell(row=i, column=j, value=v)
                c.font, c.border, c.alignment = normal, bord, haut
            for k in range(len(criteres)):
                c = f.cell(row=i, column=debut_eval + k)
                c.fill, c.border, c.alignment, c.font = fond_saisie, bord, haut, normal
        for k, (crit, choix) in enumerate(criteres):
            if not choix:
                continue
            lettre = get_column_letter(debut_eval + k)
            dv = DataValidation(type="list", formula1='"' + ",".join(choix) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Valeur non prévue",
                                error="Choisissez une valeur dans la liste.")
            dv.add(f"{lettre}3:{lettre}{derniere}")
            f.add_data_validation(dv)
        f.freeze_panes = f.cell(row=3, column=2)
        f.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}{derniere}"
        f.cell(row=1, column=1).comment = Comment("Code anonyme de la ligne : ne pas modifier.", "RépétIA")

    def niveau(it):
        return CLASSE.get(it["niveau"], it["niveau"])

    exos = [t for t in tirage if t["tache"] == "generation"]
    corrs = [t for t in tirage if t["tache"] == "correction"]
    feuille(
        "Exercices", "generation",
        [("Code", 8), ("Niveau", 13), ("Matière", 18), ("Thème demandé", 22), ("Difficulté", 10),
         ("Énoncé", 60), ("Solution", 35), ("Explication", 80)],
        [[t["code"], niveau(t["item"]), t["item"]["matiere"], t["item"]["theme"], t["item"]["difficulte"],
          t["reponse"]["enonce"].strip(), t["reponse"]["solution"].strip(), t["reponse"]["explication"].strip()]
         for t in exos],
        ["EXEMPLE", "3ème (BEPC)", "Mathématiques", "Théorème de Pythagore", "facile",
         "ABC est rectangle en A, AB = 3 cm, AC = 4 cm. Calcule BC.", "BC = 5 cm",
         "1) Le triangle est rectangle en A : on applique Pythagore. 2) BC² = 3² + 4² = 25. 3) BC = √25 = 5 cm."],
        ["Juste", "Oui", "Adapté", "5", "4", "5", "Tel quel", "Clair ; on pourrait rappeler l'hypoténuse."],
    )
    feuille(
        "Corrections", "correction",
        [("Code", 8), ("Niveau", 13), ("Matière", 18), ("Énoncé de l'exercice", 55), ("Solution attendue", 30),
         ("Réponse de l'élève", 30), ("Verdict donné", 12), ("Message à l'élève", 30), ("Explication donnée", 70)],
        [[t["code"], niveau(t["item"]), t["item"]["matiere"], _enonce_de_correction(t["item"]),
          _solution_de_correction(t["item"]), t["item"]["attendu"]["reponse_eleve"],
          "Juste" if t["reponse"]["correct"] else "Faux", t["reponse"]["verdict"].strip(),
          t["reponse"]["explication"].strip()] for t in corrs],
        ["EXEMPLE", "5ème", "Mathématiques", "Calcule (-3) + (+5).", "2", "-2", "Faux",
         "Presque ! Regardons le signe ensemble.",
         "(-3) + (+5) : les signes sont différents, on soustrait 5 - 3 = 2 et on garde le signe du plus grand, +."],
        ["Oui", "Juste", "5", "4", "5", "Tel quel", ""],
    )

    # Synthèse pour le relecteur : ses propres moyennes, recalculées par Excel.
    s = wb.create_sheet("Synthèse")
    s["A1"], s["A1"].font = "Vos notes, calculées automatiquement (ligne d'exemple exclue)", titre
    ne, nc = 2 + len(exos), 2 + len(corrs)
    lignes_synthese = [
        ("Exercices relus", f"=COUNTA(Exercices!I3:I{ne})"),
        ("Exercices justes (part)", f'=IFERROR(COUNTIF(Exercices!I3:I{ne},"Juste")/B2,"")'),
        ("Exercices utilisables tels quels (part)", f'=IFERROR(COUNTIF(Exercices!O3:O{ne},"Tel quel")/B2,"")'),
        ("Clarté moyenne", f'=IFERROR(SUMPRODUCT(VALUE(0&Exercices!L3:L{ne}))/COUNTA(Exercices!L3:L{ne}),"")'),
        ("Corrections relues", f"=COUNTA(Corrections!J3:J{nc})"),
        ("Verdicts corrects (part)", f'=IFERROR(COUNTIF(Corrections!J3:J{nc},"Oui")/B6,"")'),
        ("Corrections utilisables telles quelles (part)",
         f'=IFERROR(COUNTIF(Corrections!O3:O{nc},"Tel quel")/B6,"")'),
    ]
    for r, (lib, formule) in enumerate(lignes_synthese, 2):
        s.cell(row=r, column=1, value=lib).font = normal
        c = s.cell(row=r, column=2, value=formule)
        c.font = normal
        if "part" in lib:
            c.number_format = "0%"
        elif "moyenne" in lib:
            c.number_format = "0.0"
    s.column_dimensions["A"].width = 48
    s.column_dimensions["B"].width = 14

    # Excel et LibreOffice recalculent les formules à l'ouverture (openpyxl
    # n'enregistre pas de valeurs calculées).
    wb.calculation.fullCalcOnLoad = True
    wb.save(GRILLE)
    cle = {t["code"]: {"source": t["source"], "id": t["item"]["id"], "tache": t["tache"],
                       "niveau": t["item"]["niveau"], "matiere": t["item"]["matiere"]} for t in tirage}
    CLE.write_text(json.dumps(cle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    comptes = collections.Counter((t["tache"], t["source"]) for t in tirage)
    print(f"{len(tirage)} lignes → {GRILLE.relative_to(RACINE)} ; clé → {CLE.relative_to(RACINE)}")
    for (tache, source), n in sorted(comptes.items()):
        print(f"  {tache:<11} {source:<38} {n}")
    print("  par niveau :", dict(collections.Counter(t["item"]["niveau"] for t in tirage)))


# ---------------------------------------------------------------------------
# Analyse des retours
# ---------------------------------------------------------------------------


def analyser(fichiers: list[str]) -> None:
    from openpyxl import load_workbook

    cle = json.loads(CLE.read_text(encoding="utf-8"))
    notes: list[dict] = []
    for chemin in fichiers:
        wb = load_workbook(chemin, data_only=True)
        relecteur = wb["Consignes"]["B6"].value or pathlib.Path(chemin).stem
        for feuille, tache in (("Exercices", "generation"), ("Correction" + "s", "correction")):
            f = wb[feuille]
            entetes = [c.value for c in f[1]]
            for ligne in f.iter_rows(min_row=3, values_only=True):
                code = ligne[0]
                if code not in cle:
                    continue
                valeurs = dict(zip(entetes, ligne))
                crit = {c: valeurs.get(c) for c, _ in CRITERES[tache]}
                if not any(v not in (None, "") for k, v in crit.items() if not k.startswith("Commentaire")):
                    continue  # ligne non relue
                notes.append({"relecteur": relecteur, "code": code, **cle[code], "criteres": crit})

    def resume(lot: list[dict], tache: str) -> dict:
        if not lot:
            return {"n": 0}
        criteres = CRITERES[tache]
        sortie = {"n": len(lot)}
        chiffrees = []
        for c, choix in criteres:
            vals = [n["criteres"][c] for n in lot if n["criteres"].get(c) not in (None, "")]
            if choix == NOTE:
                nombres = [int(v) for v in vals if str(v).isdigit()]
                sortie[c] = round(statistics.mean(nombres), 2) if nombres else None
                chiffrees += nombres
            elif choix:
                sortie[c] = {k: round(v / len(vals), 3) for k, v in collections.Counter(vals).items()} if vals else None
        # Critère de la phase 4 : moyenne des notes de 1 à 5 ≥ 3,5.
        sortie["note_moyenne"] = round(statistics.mean(chiffrees), 2) if chiffrees else None
        return sortie

    rapport: dict = {"relecteurs": sorted({n["relecteur"] for n in notes}), "lignes_notees": len(notes),
                     "par_source": {}}
    for tache in ("generation", "correction"):
        for source in (MODELE, TEMOIN):
            lot = [n for n in notes if n["tache"] == tache and n["source"] == source]
            bloc = resume(lot, tache)
            bloc["par_niveau"] = {niv: resume([n for n in lot if n["niveau"] == niv], tache).get("note_moyenne")
                                  for niv in NIVEAUX if any(n["niveau"] == niv for n in lot)}
            rapport["par_source"][f"{tache} — {source}"] = bloc

    # Accord entre relecteurs sur les codes notés plusieurs fois.
    par_code = collections.defaultdict(list)
    for n in notes:
        cj = "Justesse (solution et explication)" if n["tache"] == "generation" else "Justesse de l'explication"
        if n["criteres"].get(cj):
            par_code[n["code"]].append(n["criteres"][cj])
    doubles = {c: v for c, v in par_code.items() if len(v) >= 2}
    rapport["accord_justesse"] = {
        "codes_notes_plusieurs_fois": len(doubles),
        "accord_complet": round(sum(len(set(v)) == 1 for v in doubles.values()) / len(doubles), 3) if doubles else None,
    }
    RAPPORT.write_text(json.dumps(rapport, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(rapport, ensure_ascii=False, indent=2))


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sous = p.add_subparsers(dest="action", required=True)
    sous.add_parser("preparer")
    a = sous.add_parser("analyser")
    a.add_argument("fichiers", nargs="+")
    args = p.parse_args()
    preparer() if args.action == "preparer" else analyser(args.fichiers)


if __name__ == "__main__":
    main()
